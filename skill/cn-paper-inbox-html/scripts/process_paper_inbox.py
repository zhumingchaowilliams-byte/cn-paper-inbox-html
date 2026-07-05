#!/usr/bin/env python3
"""Paper inbox processor for Chinese deep-reading Obsidian notes."""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import os
import re
import shutil
import subprocess
import sys
import time
import traceback
import urllib.request
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any


DEFAULT_INBOX = Path(os.environ.get("PAPER_INBOX_DIR", str(Path.cwd() / "paper")))
DEFAULT_VAULT = Path(os.environ.get("PAPER_VAULT_DIR", str(Path.cwd() / "vault")))
CLAUDE = Path.home() / "AppData/Roaming/npm/claude.cmd"
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}


@dataclass
class PaperRecord:
    folder: str
    folder_path: str
    status: str
    message: str
    doi: str = ""
    doi_safe: str = ""
    title: str = ""
    pages: int = 0
    pdf: str = ""
    supplements: list[str] | None = None
    images: list[str] | None = None
    missing_figures: list[str] | None = None
    image_review: list[str] | None = None
    note_path: str = ""
    html_path: str = ""
    packet_path: str = ""
    auto_image_dir: str = ""
    updated_at: str = ""


def now_iso() -> str:
    return time.strftime("%Y-%m-%dT%H:%M:%S%z")


def safe_name(text: str, max_len: int = 120) -> str:
    text = text.replace("/", "_")
    text = re.sub(r'[\\:*?"<>|]+', "_", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text[:max_len].strip(" ._") or "untitled"


def doi_to_safe(doi: str) -> str:
    return safe_name(doi.replace("/", "_"))


def read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def write_json(path: Path, data: dict[str, Any]) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def find_pdf(folder: Path) -> Path | None:
    preferred = folder / "正文.pdf"
    if preferred.exists():
        return preferred
    pdfs = sorted(folder.glob("*.pdf"))
    return pdfs[0] if pdfs else None


def find_supplements(folder: Path, main_pdf: Path | None = None) -> list[Path]:
    supplement_exts = {".docx", ".pdf", ".txt", ".md"}
    candidates: list[Path] = []
    for base in [folder / "supplements", folder]:
        if not base.exists():
            continue
        for path in sorted(base.iterdir(), key=lambda p: p.name):
            if not path.is_file() or path.suffix.lower() not in supplement_exts:
                continue
            if main_pdf and path.resolve() == main_pdf.resolve():
                continue
            if path.name in {"处理结果.md", "处理状态.json"}:
                continue
            if base == folder and path.suffix.lower() == ".pdf" and path.name == "正文.pdf":
                continue
            candidates.append(path)
    return candidates


def classify_images(folder: Path) -> dict[str, Path]:
    images: dict[str, Path] = {}
    for base in [folder, folder / "figures"]:
        if not base.exists():
            continue
        for path in sorted(base.iterdir(), key=lambda p: p.name):
            if not path.is_file() or path.suffix.lower() not in IMAGE_EXTENSIONS:
                continue
            stem = path.stem.strip()
            compact = re.sub(r"\s+", "", stem).lower()
            if "图文摘要" in stem or compact in {"graphicalabstract", "toc", "abstract"}:
                images["图文摘要"] = path
                continue
            match = re.match(r"^(?:图|fig\.?|figure)\s*0*(\d+)$", compact, re.I)
            if match:
                images[f"图{int(match.group(1))}"] = path
    return images


def missing_figures(images: dict[str, Path]) -> list[str]:
    nums = sorted(
        int(key[1:])
        for key in images
        if key.startswith("图") and key[1:].isdigit()
    )
    if not nums:
        return []
    present = set(nums)
    return [f"图{i}" for i in range(1, max(nums) + 1) if i not in present]


def figure_page_candidates(pdf: Path) -> list[tuple[int, list[str]]]:
    import pdfplumber

    candidates: list[tuple[int, list[str]]] = []
    pattern = re.compile(r"(?:Fig\.?|Figure|图)\s*\d+[A-Za-z]?", re.I)
    with pdfplumber.open(str(pdf)) as doc:
        for index, page in enumerate(doc.pages, start=1):
            text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
            matches = sorted(set(match.group(0) for match in pattern.finditer(text)))
            if matches:
                candidates.append((index, matches))
    return candidates


def clear_image_dir(output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    for old in output_dir.iterdir():
        if old.is_file():
            old.unlink()


def clamp_bbox(bbox: tuple[float, float, float, float], width: float, height: float) -> tuple[float, float, float, float]:
    x0, top, x1, bottom = bbox
    x0 = max(0.0, min(width, x0))
    x1 = max(0.0, min(width, x1))
    top = max(0.0, min(height, top))
    bottom = max(0.0, min(height, bottom))
    if x1 < x0:
        x0, x1 = x1, x0
    if bottom < top:
        top, bottom = bottom, top
    return x0, top, x1, bottom


def expand_bbox(
    bbox: tuple[float, float, float, float],
    width: float,
    height: float,
    margin: float = 10.0,
) -> tuple[float, float, float, float]:
    x0, top, x1, bottom = bbox
    return clamp_bbox((x0 - margin, top - margin, x1 + margin, bottom + margin), width, height)


def bbox_area(bbox: tuple[float, float, float, float]) -> float:
    x0, top, x1, bottom = bbox
    return max(0.0, x1 - x0) * max(0.0, bottom - top)


def bbox_iou(a: tuple[float, float, float, float], b: tuple[float, float, float, float]) -> float:
    ax0, at, ax1, ab = a
    bx0, bt, bx1, bb = b
    ix0, it = max(ax0, bx0), max(at, bt)
    ix1, ib = min(ax1, bx1), min(ab, bb)
    inter = bbox_area((ix0, it, ix1, ib))
    if inter <= 0:
        return 0.0
    return inter / max(bbox_area(a) + bbox_area(b) - inter, 1.0)


def dedupe_bboxes(candidates: list[dict[str, Any]], iou_threshold: float = 0.72) -> list[dict[str, Any]]:
    kept: list[dict[str, Any]] = []
    for candidate in sorted(candidates, key=lambda item: bbox_area(item["bbox"]), reverse=True):
        bbox = candidate["bbox"]
        if any(bbox_iou(bbox, item["bbox"]) >= iou_threshold for item in kept):
            continue
        kept.append(candidate)
    return sorted(kept, key=lambda item: (item["page"], item["bbox"][1], item["bbox"][0]))


def render_pdf_bbox(
    pdfium_doc: Any,
    page_index: int,
    bbox: tuple[float, float, float, float],
    dest: Path,
    scale: float = 2.5,
) -> tuple[int, int]:
    page = pdfium_doc[page_index]
    try:
        bitmap = page.render(scale=scale)
        image = bitmap.to_pil()
        x0, top, x1, bottom = bbox
        crop_box = (
            int(round(x0 * scale)),
            int(round(top * scale)),
            int(round(x1 * scale)),
            int(round(bottom * scale)),
        )
        cropped = image.crop(crop_box)
        cropped.save(dest)
        return cropped.width, cropped.height
    finally:
        close = getattr(page, "close", None)
        if callable(close):
            close()


def extract_pdf_figure_pages(pdf: Path, output_dir: Path, max_pages: int = 24) -> list[dict[str, Any]]:
    """Render pages that appear to contain figures as a last-resort fallback."""
    try:
        import pypdfium2 as pdfium
    except Exception as exc:
        raise RuntimeError("pypdfium2 is required for PDF rendering. Install with `pip install pypdfium2`.") from exc

    clear_image_dir(output_dir)

    candidates = figure_page_candidates(pdf)
    if not candidates:
        candidates = [(index, []) for index in range(1, min(max_pages, 6) + 1)]
    candidates = candidates[:max_pages]

    manifest: list[dict[str, Any]] = []
    doc = pdfium.PdfDocument(str(pdf))
    try:
        for page_number, labels in candidates:
            name = f"pdf-page-{page_number:03d}.png"
            dest = output_dir / name
            page = doc[page_number - 1]
            try:
                width, height = page.get_size()
            finally:
                close = getattr(page, "close", None)
                if callable(close):
                    close()
            out_width, out_height = render_pdf_bbox(doc, page_number - 1, (0.0, 0.0, width, height), dest, scale=2.0)
            manifest.append(
                {
                    "page": page_number,
                    "file": str(dest),
                    "width": out_width,
                    "height": out_height,
                    "labels": labels,
                    "source": "full-page-fallback",
                    "note": "PDF page rendered automatically for agent-side visual inspection.",
                }
            )
    finally:
        doc.close()
    (output_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def caption_lines(words: list[dict[str, Any]]) -> list[dict[str, Any]]:
    pattern = re.compile(r"^(?:Fig\.?|Figure|图)\s*\d+[A-Za-z]?[.:：]?$", re.I)
    lines: dict[tuple[int, int], list[dict[str, Any]]] = {}
    for word in words:
        key = (int(round(float(word.get("top", 0)) / 4)), int(round(float(word.get("bottom", 0)) / 4)))
        lines.setdefault(key, []).append(word)
    captions: list[dict[str, Any]] = []
    for grouped in lines.values():
        grouped = sorted(grouped, key=lambda w: float(w.get("x0", 0)))
        text = " ".join(str(w.get("text", "")) for w in grouped)
        if not re.search(r"\b(?:Fig\.?|Figure)\s*\d+|图\s*\d+", text, flags=re.I):
            continue
        if not any(pattern.match(str(w.get("text", ""))) for w in grouped) and not re.search(r"图\s*\d+", text):
            continue
        captions.append(
            {
                "text": text,
                "x0": min(float(w["x0"]) for w in grouped),
                "x1": max(float(w["x1"]) for w in grouped),
                "top": min(float(w["top"]) for w in grouped),
                "bottom": max(float(w["bottom"]) for w in grouped),
            }
        )
    return sorted(captions, key=lambda item: (item["top"], item["x0"]))


def object_bbox(obj: dict[str, Any]) -> tuple[float, float, float, float] | None:
    try:
        return float(obj["x0"]), float(obj["top"]), float(obj["x1"]), float(obj["bottom"])
    except Exception:
        return None


def visual_object_candidates(page: Any, min_area: float) -> list[tuple[float, float, float, float]]:
    boxes: list[tuple[float, float, float, float]] = []
    for collection in [page.images, page.rects, page.curves, page.lines]:
        for obj in collection:
            bbox = object_bbox(obj)
            if not bbox:
                continue
            if bbox_area(bbox) >= min_area:
                boxes.append(bbox)
    return boxes


def union_bboxes(boxes: list[tuple[float, float, float, float]]) -> tuple[float, float, float, float]:
    return (
        min(box[0] for box in boxes),
        min(box[1] for box in boxes),
        max(box[2] for box in boxes),
        max(box[3] for box in boxes),
    )


def caption_guided_bbox(
    caption: dict[str, Any],
    visual_boxes: list[tuple[float, float, float, float]],
    page_width: float,
    page_height: float,
) -> tuple[float, float, float, float] | None:
    caption_top = float(caption["top"])
    caption_x0 = float(caption["x0"])
    caption_x1 = float(caption["x1"])
    caption_width = caption_x1 - caption_x0
    full_width_caption = caption_width > page_width * 0.42
    if full_width_caption:
        x0, x1 = 0.0, page_width
    else:
        center = (caption_x0 + caption_x1) / 2
        if center < page_width / 2:
            x0, x1 = 0.0, page_width / 2 + 18
        else:
            x0, x1 = page_width / 2 - 18, page_width
    window_top = max(0.0, caption_top - page_height * 0.62)
    nearby = [
        box
        for box in visual_boxes
        if box[3] <= caption_top + 8
        and box[1] >= window_top
        and not (box[2] < x0 or box[0] > x1)
    ]
    if nearby:
        return expand_bbox(union_bboxes(nearby), page_width, page_height, margin=12)

    crop_top = max(0.0, caption_top - page_height * 0.48)
    crop_bottom = max(crop_top + 80, caption_top - 4)
    bbox = clamp_bbox((x0, crop_top, x1, crop_bottom), page_width, page_height)
    return bbox if bbox_area(bbox) >= 45000 else None


def extract_pdf_figure_regions(
    pdf: Path,
    output_dir: Path,
    min_width: int = 160,
    min_height: int = 120,
    min_area: int = 45000,
    max_regions: int = 80,
) -> list[dict[str, Any]]:
    """Crop likely figure regions using pdfplumber coordinates and pypdfium2 rendering."""
    try:
        import pdfplumber
        import pypdfium2 as pdfium
    except Exception as exc:
        raise RuntimeError("pdfplumber and pypdfium2 are required for figure-region extraction.") from exc

    clear_image_dir(output_dir)

    manifest: list[dict[str, Any]] = []
    candidates: list[dict[str, Any]] = []
    with pdfplumber.open(str(pdf)) as plumber_doc:
        for page_index, page in enumerate(plumber_doc.pages):
            page_width, page_height = float(page.width), float(page.height)
            visual_boxes = visual_object_candidates(page, min_area=2500)
            for image_index, bbox in enumerate(page.images, start=1):
                obj_bbox = object_bbox(bbox)
                if not obj_bbox:
                    continue
                if bbox_area(obj_bbox) < min_area:
                    continue
                candidates.append(
                    {
                        "page": page_index + 1,
                        "bbox": expand_bbox(obj_bbox, page_width, page_height, margin=8),
                        "source": "pdfplumber-image-object",
                        "label": f"image-object-{image_index}",
                    }
                )
            try:
                words = page.extract_words(x_tolerance=1, y_tolerance=3) or []
            except Exception:
                words = []
            for caption_index, caption in enumerate(caption_lines(words), start=1):
                bbox = caption_guided_bbox(caption, visual_boxes, page_width, page_height)
                if bbox and bbox_area(bbox) >= min_area:
                    candidates.append(
                        {
                            "page": page_index + 1,
                            "bbox": bbox,
                            "source": "caption-guided-region",
                            "label": caption["text"][:120],
                            "caption": caption["text"],
                        }
                    )

    candidates = dedupe_bboxes(candidates)[:max_regions]
    pdfium_doc = pdfium.PdfDocument(str(pdf))
    try:
        for index, item in enumerate(candidates, start=1):
            page = int(item["page"])
            bbox = item["bbox"]
            name = f"pdf-figure-p{page:03d}-{index:02d}.png"
            dest = output_dir / name
            width, height = render_pdf_bbox(pdfium_doc, page - 1, bbox, dest, scale=2.8)
            if width < min_width or height < min_height:
                dest.unlink(missing_ok=True)
                continue
            digest = hashlib.sha256(dest.read_bytes()).hexdigest()
            if any(entry.get("sha256") == digest for entry in manifest):
                dest.unlink(missing_ok=True)
                continue
            x0, top, x1, bottom = bbox
            manifest.append(
                {
                    "page": page,
                    "file": str(dest),
                    "width": width,
                    "height": height,
                    "bbox": {
                        "x0": round(x0, 2),
                        "top": round(top, 2),
                        "x1": round(x1, 2),
                        "bottom": round(bottom, 2),
                    },
                    "source": item["source"],
                    "label": item.get("label", ""),
                    "caption": item.get("caption", ""),
                    "sha256": digest,
                    "note": "Figure region cropped from PDF coordinates via pdfplumber + pypdfium2, not a full-page screenshot.",
                }
            )
    finally:
        close = getattr(pdfium_doc, "close", None)
        if callable(close):
            close()
    (output_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def extract_pdf_embedded_images(
    pdf: Path,
    output_dir: Path,
    min_width: int = 160,
    min_height: int = 120,
    min_area: int = 45000,
    max_images: int = 80,
) -> list[dict[str, Any]]:
    """Compatibility wrapper: crop figure-like regions before falling back to full-page renders."""
    return extract_pdf_figure_regions(
        pdf=pdf,
        output_dir=output_dir,
        min_width=min_width,
        min_height=min_height,
        min_area=min_area,
        max_regions=max_images,
    )


def extract_pdf_metadata(pdf: Path) -> tuple[int, str, str, str]:
    from pypdf import PdfReader

    reader = PdfReader(str(pdf))
    pages = len(reader.pages)
    meta_title = (reader.metadata.title or "").strip() if reader.metadata else ""
    first_text_parts = []
    for index in range(min(4, pages)):
        try:
            first_text_parts.append(reader.pages[index].extract_text() or "")
        except Exception:
            pass
    first_text = "\n".join(first_text_parts)
    doi_match = re.search(r"\b10\.\d{4,9}/[-._;()/:A-Z0-9]+", first_text, re.I)
    doi = doi_match.group(0).rstrip(".,;") if doi_match else ""
    title = meta_title or infer_title_from_text(first_text)
    return pages, title, doi, first_text


def infer_title_from_text(text: str) -> str:
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines()]
    lines = [line for line in lines if line and len(line) > 12]
    if not lines:
        return ""
    return lines[0][:220]


def extract_supplement_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        parts: list[str] = []
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
        for table_index, table in enumerate(doc.tables, start=1):
            rows = []
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                rows.append(" | ".join(cells))
            if rows:
                parts.append(f"[Supplement Table {table_index}]\n" + "\n".join(rows))
        return "\n\n".join(parts)
    if suffix == ".pdf":
        import pdfplumber

        chunks: list[str] = []
        with pdfplumber.open(str(path)) as doc:
            for index, page in enumerate(doc.pages, start=1):
                text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
                chunks.append(f"### Supplement Page {index}\n{text.strip()}")
        return "\n\n".join(chunks)
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace")
    return ""


def extract_supplements_packet(supplements: list[Path]) -> str:
    sections: list[str] = []
    for path in supplements:
        try:
            text = extract_supplement_text(path).strip()
        except Exception as exc:
            text = f"[supplement extraction failed: {exc}]"
        if text:
            sections.append(f"## Supplementary Source: {path.name}\n\n{text}")
    return "\n\n".join(sections)


def extract_pdf_packet(pdf: Path, output: Path, title: str, doi: str, supplements: list[Path] | None = None) -> None:
    import pdfplumber

    chunks: list[str] = []
    with pdfplumber.open(str(pdf)) as doc:
        for index, page in enumerate(doc.pages, start=1):
            try:
                page_text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
            except Exception as exc:
                page_text = f"[page text extraction failed: {exc}]"
            chunks.append(f"## Page {index}\n\n{page_text.strip()}")
    body = "\n\n".join(chunks)
    supplement_body = extract_supplements_packet(supplements or [])
    if not supplement_body:
        supplement_body = "未提供补充材料。"
    output.write_text(
        f"# Paper Text Packet\n\n- Source PDF: `{pdf}`\n- Title: {title}\n- DOI: {doi}\n- Supplement count: {len(supplements or [])}\n\n# Main Paper\n\n{body}\n\n# Supplementary Materials\n\n{supplement_body}\n",
        encoding="utf-8",
    )


def prompt_text(vault: Path) -> str:
    output_rules = (
        "\n\n输出硬性要求：\n"
        "- 直接从 Markdown 一级标题开始输出，例如 `# 中文题目`。\n"
        "- 不要写任何客套话、确认语、说明语或前言，例如“好的，遵照您的指示”“以下是”。\n"
        "- 不要输出 YAML frontmatter、笔记属性、代码围栏或元说明。\n"
        "- 如果没有图片文件，也要保留图文导读，但只根据论文文本和图题解释，不虚构视觉细节。\n"
    )
    candidates = sorted((vault / "Prompts").glob("*.md"), key=lambda p: p.stat().st_mtime, reverse=True)
    for path in candidates:
        text = path.read_text(encoding="utf-8", errors="replace")
        if "中文深读" in text and "实验知识库" in text:
            return text + "\n\n补充要求：如果文本包包含 Supplementary Materials，必须优先从补充材料中补充材料合成方法、表征细节、补充图表、实验条件和对照组。若未提供补充材料，请在人工核查清单中提醒。" + output_rules
    return """你是一名环境材料与水处理方向的中文论文精读助手。请输出中文深读笔记，包含：题目、关键词、一句话结论、文章信息、成果简介、全文速览、图文导读、方法细读、关键实验卡片、机制链条、我以后可以怎么用、人工核查清单。所有关键结论必须绑定图号或表号；没有说明的实验条件写“文本未说明”。如果文本包包含 Supplementary Materials，必须优先从补充材料中补充材料合成方法、表征细节、补充图表、实验条件和对照组。""" + output_rules


def call_deepseek(prompt: str, source_text: str, model: str = "deepseek-v4-flash") -> str:
    api_key = os.environ.get("DEEPSEEK_API_KEY")
    if not api_key:
        raise RuntimeError("DEEPSEEK_API_KEY is not configured")
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": prompt},
            {"role": "user", "content": "以下是论文文本包，请按要求输出中文深读笔记：\n\n" + source_text},
        ],
        "stream": False,
    }
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(
        "https://api.deepseek.com/chat/completions",
        data=data,
        headers={"Content-Type": "application/json", "Authorization": f"Bearer {api_key}"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=1200) as response:
        result = json.loads(response.read().decode("utf-8"))
    return result["choices"][0]["message"]["content"]


def call_claude(prompt: str, source_text: str, budget: str = "1") -> str:
    if not CLAUDE.exists():
        raise RuntimeError(f"Claude CLI not found: {CLAUDE}")
    timeout_seconds = int(os.environ.get("PAPER_INBOX_MODEL_TIMEOUT", "3600"))
    try:
        proc = subprocess.run(
            [
                str(CLAUDE),
                "-p",
                "--tools",
                "",
                "--no-session-persistence",
                "--max-budget-usd",
                budget,
                "--model",
                "sonnet",
            ],
            input=prompt + "\n\n以下是论文文本包，请按上述协议输出：\n\n" + source_text,
            text=True,
            encoding="utf-8",
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=timeout_seconds,
        )
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(f"Claude CLI 超时：超过 {timeout_seconds} 秒仍未返回。可点“重新处理”重试，或设置 DEEPSEEK_API_KEY 后走 DeepSeek。") from exc
    if proc.returncode != 0:
        details = (proc.stderr or proc.stdout or "").strip()
        if len(details) > 4000:
            details = details[-4000:]
        raise RuntimeError(f"Claude CLI 返回错误码 {proc.returncode}。\n{details}")
    return proc.stdout.strip()


def generate_deep_reading(vault: Path, packet: Path, model_route: str = "auto") -> tuple[str, str]:
    prompt = prompt_text(vault)
    source_text = packet.read_text(encoding="utf-8", errors="replace")
    if model_route in {"auto", "deepseek"} and os.environ.get("DEEPSEEK_API_KEY"):
        return call_deepseek(prompt, source_text), "deepseek"
    if model_route == "deepseek":
        raise RuntimeError("DEEPSEEK_API_KEY is not configured")
    return call_claude(prompt, source_text), "claude"


def copy_assets(images: dict[str, Path], target_dir: Path) -> dict[str, str]:
    if target_dir.exists():
        for old in target_dir.iterdir():
            if old.is_file():
                old.unlink()
    target_dir.mkdir(parents=True, exist_ok=True)
    copied: dict[str, str] = {}
    for label, src in sorted(images.items(), key=lambda item: image_sort_key(item[0])):
        ext = src.suffix.lower()
        dest = target_dir / f"{label}{ext}"
        shutil.copy2(src, dest)
        copied[label] = dest.name
    return copied


def image_sort_key(label: str) -> tuple[int, int]:
    if label == "图文摘要":
        return (0, 0)
    match = re.match(r"图(\d+)$", label)
    return (1, int(match.group(1))) if match else (2, 999)


def insert_images(markdown_text: str, asset_rel_dir: str, copied: dict[str, str]) -> tuple[str, list[str]]:
    text = markdown_text.strip() + "\n"
    review: list[str] = []
    graphical = copied.get("图文摘要")
    if graphical and f"{asset_rel_dir}/{graphical}" not in text:
        embed = f"\n![[{asset_rel_dir}/{graphical}]]\n"
        if "## 成果简介" in text:
            text = text.replace("## 成果简介", embed + "\n## 成果简介", 1)
        else:
            text += "\n## 图文摘要\n" + embed

    for label, filename in sorted(copied.items(), key=lambda item: image_sort_key(item[0])):
        match = re.match(r"图(\d+)$", label)
        if not match:
            continue
        num = int(match.group(1))
        if f"{asset_rel_dir}/{filename}" in text:
            continue
        embed = f"\n![[{asset_rel_dir}/{filename}]]\n"
        patterns = [
            rf"(###\s*Fig\.\s*{num}\b[^\n]*\n)",
            rf"(###\s*图\s*{num}\b[^\n]*\n)",
            rf"(###\s*Figure\s*{num}\b[^\n]*\n)",
            rf"(\*\*Fig\.\s*{num}\b[^\n]*\*\*[^\n]*\n)",
            rf"(\*\*图\s*{num}\b[^\n]*\*\*[^\n]*\n)",
            rf"(^Fig\.\s*{num}\b[^\n]*\n)",
            rf"(^图\s*{num}\b[^\n]*\n)",
        ]
        inserted = False
        for pattern in patterns:
            text, count = re.subn(pattern, rf"\1{embed}\n", text, count=1, flags=re.I)
            if count:
                inserted = True
                break
        if not inserted:
            appendix = "\n## 附图\n" if "## 附图" not in text else "\n"
            text += f"{appendix}\n### {label}\n{embed}\n"
            review.append(f"{label} 未匹配到对应 Fig. 小节，已放入附图")
    return text, review


def clean_model_output(markdown_text: str) -> str:
    text = markdown_text.strip()
    text = re.sub(r"^```(?:markdown|md)?\s*", "", text, flags=re.I).strip()
    text = re.sub(r"\s*```$", "", text).strip()
    text = re.sub(r"^---\s*\n.*?\n---\s*\n", "", text, flags=re.S).strip()
    first_heading = re.search(r"(?m)^#\s+", text)
    if first_heading:
        text = text[first_heading.start() :].strip()
    lines = text.splitlines()
    while lines and re.match(r"^\s*(好的|遵照|以下是|根据提供|我将|下面是)", lines[0]):
        lines.pop(0)
    return "\n".join(lines).strip() + "\n"


def extract_generated_title(markdown_text: str, fallback: str) -> str:
    lines = [line.strip() for line in markdown_text.splitlines()]
    for index, line in enumerate(lines):
        if line in {"# 题目", "# 标题"}:
            for candidate in lines[index + 1 : index + 8]:
                if candidate and not candidate.startswith("#"):
                    return candidate.strip("* ") or fallback
        line = line.strip()
        if line.startswith("# "):
            title = re.sub(r"^#\s+", "", line).strip("* ")
            if title not in {"题目", "标题"}:
                return title or fallback
    return fallback


def unique_output_path(output_dir: Path, title: str, doi_safe: str, suffix: str, existing: str = "") -> Path:
    base = safe_name(title, max_len=90)
    candidate = output_dir / f"{base}{suffix}"
    if existing and Path(existing) == candidate:
        return candidate
    if not candidate.exists():
        return candidate
    candidate = output_dir / f"{base} - {doi_safe}{suffix}"
    if not candidate.exists():
        return candidate
    index = 2
    while True:
        numbered = output_dir / f"{base} - {doi_safe} ({index}){suffix}"
        if not numbered.exists():
            return numbered
        index += 1


def markdown_to_html(markdown_text: str, title: str, html_path: Path, vault: Path) -> str:
    def convert_obsidian_image(match: re.Match[str]) -> str:
        target = match.group(1).strip()
        src_path = vault / target.replace("/", os.sep)
        try:
            src = os.path.relpath(src_path, html_path.parent).replace("\\", "/")
        except ValueError:
            src = src_path.as_uri()
        alt = Path(target).stem
        return f'<figure><img src="{html.escape(src)}" alt="{html.escape(alt)}"><figcaption>{html.escape(alt)}</figcaption></figure>'

    converted = re.sub(r"!\[\[([^\]]+)\]\]", convert_obsidian_image, markdown_text)
    try:
        import markdown  # type: ignore

        body = markdown.markdown(converted, extensions=["tables", "fenced_code", "toc"])
    except Exception:
        body = fallback_markdown_to_html(converted)
    return html_document(title, body)


def fallback_markdown_to_html(markdown_text: str) -> str:
    lines = markdown_text.splitlines()
    parts: list[str] = []
    in_list = False
    paragraph: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph
        if paragraph:
            parts.append("<p>" + html.escape(" ".join(paragraph)) + "</p>")
            paragraph = []

    def close_list() -> None:
        nonlocal in_list
        if in_list:
            parts.append("</ul>")
            in_list = False

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            close_list()
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            close_list()
            level = len(heading.group(1))
            parts.append(f"<h{level}>{html.escape(heading.group(2))}</h{level}>")
            continue
        if stripped.startswith("<figure>"):
            flush_paragraph()
            close_list()
            parts.append(stripped)
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        if bullet:
            flush_paragraph()
            if not in_list:
                parts.append("<ul>")
                in_list = True
            parts.append(f"<li>{html.escape(bullet.group(1))}</li>")
            continue
        paragraph.append(stripped)
    flush_paragraph()
    close_list()
    return "\n".join(parts)


def html_document(title: str, body: str) -> str:
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{html.escape(title)}</title>
  <style>
    :root {{
      color-scheme: light;
      --ink: #17202a;
      --muted: #5d6875;
      --line: #d8dee6;
      --paper: #ffffff;
      --soft: #f6f8fa;
      --accent: #166a5f;
    }}
    body {{
      margin: 0;
      background: #eef2f5;
      color: var(--ink);
      font-family: "Microsoft YaHei", "PingFang SC", "Noto Sans CJK SC", Arial, sans-serif;
      line-height: 1.74;
    }}
    main {{
      max-width: 980px;
      margin: 28px auto;
      padding: 42px 52px;
      background: var(--paper);
      box-shadow: 0 10px 32px rgba(23, 32, 42, 0.08);
    }}
    h1 {{
      font-size: 30px;
      line-height: 1.28;
      margin: 0 0 22px;
      color: #102a3a;
      border-bottom: 3px solid var(--accent);
      padding-bottom: 16px;
    }}
    h2 {{
      font-size: 22px;
      margin: 34px 0 14px;
      padding-left: 12px;
      border-left: 5px solid var(--accent);
      color: #193645;
    }}
    h3 {{
      font-size: 18px;
      margin: 24px 0 10px;
      color: #23384a;
    }}
    p, li {{ font-size: 16px; }}
    ul, ol {{ padding-left: 1.4em; }}
    blockquote {{
      margin: 18px 0;
      padding: 12px 18px;
      background: var(--soft);
      border-left: 4px solid var(--line);
      color: var(--muted);
    }}
    table {{
      width: 100%;
      border-collapse: collapse;
      margin: 18px 0;
      font-size: 14px;
    }}
    th, td {{
      border: 1px solid var(--line);
      padding: 8px 10px;
      vertical-align: top;
    }}
    th {{ background: var(--soft); }}
    code {{
      background: var(--soft);
      padding: 1px 5px;
      border-radius: 4px;
    }}
    figure {{
      margin: 24px 0;
      padding: 14px;
      background: var(--soft);
      border: 1px solid var(--line);
    }}
    figure img {{
      display: block;
      max-width: 100%;
      height: auto;
      margin: 0 auto;
    }}
    figcaption {{
      margin-top: 8px;
      text-align: center;
      color: var(--muted);
      font-size: 14px;
    }}
    @media (max-width: 720px) {{
      main {{ margin: 0; padding: 24px 18px; }}
      h1 {{ font-size: 24px; }}
      h2 {{ font-size: 20px; }}
    }}
  </style>
</head>
<body>
<main>
{body}
</main>
</body>
</html>
"""


def write_html_output(markdown_text: str, html_path: Path, title: str, vault: Path) -> None:
    html_path.write_text(markdown_to_html(markdown_text, title, html_path, vault), encoding="utf-8")


def unique_note_path(note_dir: Path, title: str, doi_safe: str, existing: str = "") -> Path:
    return unique_output_path(note_dir, title, doi_safe, ".md", existing=existing)


def unique_html_path(html_dir: Path, title: str, doi_safe: str, existing: str = "") -> Path:
    return unique_output_path(html_dir, title, doi_safe, ".html", existing=existing)


def legacy_frontmatter(doi: str, title: str, source_model: str) -> str:
    return "\n".join(
        [
            "---",
            "type: paper_deep_reading",
            f'source: "{source_model}"',
            f'doi: "{doi}"',
            f'title: "{title.replace(chr(34), chr(39))}"',
            "status: ai_extracted_unchecked",
            "layout: polished",
            "---",
            "",
        ]
    )


def scan_folder(folder: Path, vault: Path = DEFAULT_VAULT) -> PaperRecord:
    status_data = read_json(folder / "处理状态.json")
    pdf = find_pdf(folder)
    supplements = find_supplements(folder, pdf)
    images = classify_images(folder)
    missing = missing_figures(images)
    if not pdf:
        return PaperRecord(
            folder=folder.name,
            folder_path=str(folder),
            status="needs_retry",
            message="缺少 正文.pdf 或 PDF 文件",
            supplements=[p.name for p in supplements],
            images=list(images.keys()),
            missing_figures=missing,
            image_review=[],
            updated_at=now_iso(),
        )
    try:
        pages, title, doi, _ = extract_pdf_metadata(pdf)
    except Exception as exc:
        return PaperRecord(
            folder=folder.name,
            folder_path=str(folder),
            status="needs_retry",
            message=f"PDF 读取失败: {exc}",
            pdf=str(pdf),
            supplements=[p.name for p in supplements],
            images=list(images.keys()),
            missing_figures=missing,
            image_review=[],
            updated_at=now_iso(),
        )
    doi_safe = doi_to_safe(doi) if doi else safe_name(folder.name)
    stored_note = status_data.get("note_path", "")
    stored_html = status_data.get("html_path", "")
    stored_packet = status_data.get("packet_path", "")
    stored_auto_image_dir = status_data.get("auto_image_dir", "")
    note_path = Path(stored_note) if stored_note else vault / "Knowledge/Paper Deep Readings" / f"{doi_safe} - polished.md"
    html_path = Path(stored_html) if stored_html else vault / "Knowledge/Paper Deep Readings HTML" / f"{doi_safe} - polished.html"
    packet_path = Path(stored_packet) if stored_packet else vault / "AI Inputs" / f"{doi_safe}.paper_text_packet.md"
    auto_image_dir = Path(stored_auto_image_dir) if stored_auto_image_dir else vault / "Assets/Papers" / doi_safe / "pdf-auto-pages"
    status = status_data.get("status") or ("done" if note_path.exists() else "new")
    if not doi and status == "new":
        status = "needs_review"
    message = status_data.get("message") or ("缺少 DOI，需人工复核" if not doi else "可处理")
    if missing:
        message = f"{message}；缺图提示: {', '.join(missing)}"
    return PaperRecord(
        folder=folder.name,
        folder_path=str(folder),
        status=status,
        message=message,
        doi=doi,
        doi_safe=doi_safe,
        title=title,
        pages=pages,
        pdf=str(pdf),
        supplements=[p.name for p in supplements],
        images=list(images.keys()),
        missing_figures=missing,
        image_review=status_data.get("image_review", []),
        note_path=str(note_path) if note_path.exists() else "",
        html_path=str(html_path) if html_path.exists() else "",
        packet_path=str(packet_path) if packet_path.exists() else "",
        auto_image_dir=str(auto_image_dir) if auto_image_dir.exists() else "",
        updated_at=status_data.get("updated_at", now_iso()),
    )


def scan_inbox(inbox: Path = DEFAULT_INBOX, vault: Path = DEFAULT_VAULT) -> list[PaperRecord]:
    if not inbox.exists():
        return []
    folders = [path for path in sorted(inbox.iterdir(), key=lambda p: p.name) if path.is_dir()]
    return [scan_folder(folder, vault=vault) for folder in folders]


def prepare_folder(folder: Path, vault: Path = DEFAULT_VAULT, extract_auto_images: bool = True) -> PaperRecord:
    """Prepare deterministic inputs for an external Codex/Claude agent without calling a model API."""
    record = scan_folder(folder, vault=vault)
    status_path = folder / "处理状态.json"
    if not record.pdf:
        record.status = "needs_retry"
        record.message = "缺少 正文.pdf 或 PDF 文件"
        record.updated_at = now_iso()
        write_json(status_path, asdict(record))
        return record
    if not record.doi:
        record.status = "needs_review"
        record.message = "未识别 DOI；仍可人工继续，但引用信息需复核"

    doi_safe = record.doi_safe or safe_name(folder.name)
    input_dir = vault / "AI Inputs"
    asset_dir = vault / "Assets/Papers" / doi_safe
    extracted_image_dir = asset_dir / "pdf-extracted-images"
    auto_image_dir = asset_dir / "pdf-auto-pages"
    input_dir.mkdir(parents=True, exist_ok=True)
    asset_dir.mkdir(parents=True, exist_ok=True)
    packet_path = input_dir / f"{doi_safe}.paper_text_packet.md"
    supplements = find_supplements(folder, Path(record.pdf))
    extract_pdf_packet(Path(record.pdf), packet_path, record.title, record.doi, supplements=supplements)

    manual_images = classify_images(folder)
    if manual_images:
        copy_assets(manual_images, asset_dir)
    auto_manifest: list[dict[str, Any]] = []
    auto_warning = ""
    visual_dir = Path("")
    visual_mode = "none"
    if extract_auto_images and not manual_images:
        try:
            auto_manifest = extract_pdf_embedded_images(Path(record.pdf), extracted_image_dir)
            if auto_manifest:
                visual_dir = extracted_image_dir
                visual_mode = "embedded-images"
            else:
                auto_manifest = extract_pdf_figure_pages(Path(record.pdf), auto_image_dir)
                if auto_manifest:
                    visual_dir = auto_image_dir
                    visual_mode = "page-renders"
        except Exception as exc:
            auto_warning = f"；自动抽图失败: {type(exc).__name__}: {exc}"

    agent_task_path = input_dir / f"{doi_safe}.agent_task.md"
    agent_task_path.write_text(
        "\n".join(
            [
                "# Agent Task",
                "",
                "请使用当前 Codex/Claude Agent 直接生成中文深读 Markdown，不要调用 DeepSeek API。",
                "",
                f"- Text packet: `{packet_path}`",
                f"- Auto visual assets: `{visual_dir}`" if auto_manifest else "- Auto visual assets: 未生成",
                f"- Auto visual mode: {visual_mode}",
                "- Output must start directly with `# 中文题目`.",
                "- Do not include YAML/frontmatter/note properties.",
                "- Do not include preambles such as `好的，遵照您的指示` or `以下是`.",
                "- If images are available, inspect them directly when your runtime supports image viewing; otherwise rely on captions and text.",
                "- Do not make a text-only `图文导读` when visual assets exist.",
                "- Insert each available figure under its matching `Fig. N` / `图N` heading with Obsidian embeds such as `![[Assets/Papers/<doi-safe>/图1.png]]`, then write the Chinese interpretation below it.",
                "- If only full-page renders are available, crop the figure region when practical; otherwise embed the page render and label it as a fallback figure page.",
                "",
                "After writing the generated Markdown to a temporary file, run:",
                "",
                "```bash",
                f"python scripts/process_paper_inbox.py --inbox <inbox> --vault <vault> --finalize {folder.name} --generated-md <generated-note.md>",
                "```",
                "",
            ]
        ),
        encoding="utf-8",
    )

    prepared = scan_folder(folder, vault=vault)
    prepared.status = "prepared" if record.doi else "needs_review"
    prepared.message = "已准备文本包" + (
        f"；自动抽取视觉素材 {len(auto_manifest)} 张（{visual_mode}）" if auto_manifest else "；使用手动图片或未抽取图片"
    ) + auto_warning
    prepared.packet_path = str(packet_path)
    prepared.auto_image_dir = str(visual_dir) if auto_manifest else ""
    prepared.supplements = [p.name for p in supplements]
    prepared.updated_at = now_iso()
    write_json(status_path, asdict(prepared))
    return prepared


def append_auto_images(markdown_text: str, auto_image_dir: Path, vault: Path) -> str:
    manifest_path = auto_image_dir / "manifest.json"
    if not manifest_path.exists():
        return markdown_text
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except Exception:
        return markdown_text
    if not manifest:
        return markdown_text
    text = markdown_text.rstrip() + "\n"
    section_title = "## PDF 自动抽取图片"
    if "pdf-auto-pages" in str(auto_image_dir):
        section_title = "## PDF 自动抽取图页"
    if section_title not in text:
        text += f"\n{section_title}\n"
    for item in manifest:
        file_path = Path(item.get("file", ""))
        if not file_path.exists():
            continue
        try:
            rel = os.path.relpath(file_path, vault).replace("\\", "/")
        except ValueError:
            rel = str(file_path)
        if rel in text:
            continue
        labels = ", ".join(item.get("labels") or [])
        source = item.get("source", "")
        if source == "embedded-image":
            caption = f"Page {item.get('page')} extracted image"
            if item.get("width") and item.get("height"):
                caption += f"；{item.get('width')} x {item.get('height')}"
        else:
            caption = f"Page {item.get('page')}" + (f"；检测到：{labels}" if labels else "")
        text += f"\n### {caption}\n![[{rel}]]\n"
    return text


def finalize_folder(folder: Path, generated_md: Path, vault: Path = DEFAULT_VAULT) -> PaperRecord:
    """Turn an agent-generated Markdown note into cleaned Markdown plus HTML."""
    record = scan_folder(folder, vault=vault)
    status_path = folder / "处理状态.json"
    if not record.pdf:
        record.status = "needs_retry"
        record.message = "无法 finalize：缺少 正文.pdf 或 PDF 文件"
        write_json(status_path, asdict(record))
        return record

    doi_safe = record.doi_safe or safe_name(folder.name)
    note_dir = vault / "Knowledge/Paper Deep Readings"
    html_dir = vault / "Knowledge/Paper Deep Readings HTML"
    asset_dir = vault / "Assets/Papers" / doi_safe
    note_dir.mkdir(parents=True, exist_ok=True)
    html_dir.mkdir(parents=True, exist_ok=True)
    asset_dir.mkdir(parents=True, exist_ok=True)

    raw = generated_md.read_text(encoding="utf-8", errors="replace")
    note_body = clean_model_output(raw)
    manual_images = classify_images(folder)
    image_review: list[str] = []
    if manual_images:
        copied = copy_assets(manual_images, asset_dir)
        note_body, image_review = insert_images(note_body, f"Assets/Papers/{doi_safe}", copied)
    else:
        extracted_dir = asset_dir / "pdf-extracted-images"
        page_dir = asset_dir / "pdf-auto-pages"
        note_body = append_auto_images(note_body, extracted_dir, vault)
        note_body = append_auto_images(note_body, page_dir, vault)
    note_body = clean_model_output(note_body)
    generated_title = extract_generated_title(note_body, record.title or folder.name)

    existing = read_json(status_path)
    note_path = unique_note_path(note_dir, generated_title, doi_safe, existing=existing.get("note_path", ""))
    html_path = unique_html_path(html_dir, generated_title, doi_safe, existing=existing.get("html_path", ""))
    note_path.write_text(note_body, encoding="utf-8")
    write_html_output(note_body, html_path, generated_title, vault)

    final_record = scan_folder(folder, vault=vault)
    final_record.status = "done" if not image_review else "needs_review"
    final_record.message = "已由当前 Agent 生成 Markdown 和 HTML" + (
        f"；图片需复核: {'; '.join(image_review)}" if image_review else ""
    )
    final_record.note_path = str(note_path)
    final_record.html_path = str(html_path)
    final_record.image_review = image_review
    final_record.packet_path = existing.get("packet_path", final_record.packet_path)
    final_record.auto_image_dir = existing.get("auto_image_dir", final_record.auto_image_dir)
    final_record.updated_at = now_iso()
    write_json(status_path, asdict(final_record))
    result_path = folder / "处理结果.md"
    result_path.write_text(
        "\n".join(
            [
                "# 处理结果",
                "",
                f"- 状态：{final_record.status}",
                f"- DOI：{final_record.doi or '未识别'}",
                f"- 标题：{final_record.title}",
                f"- Markdown：{note_path}",
                f"- HTML：{html_path}",
                f"- 更新时间：{final_record.updated_at}",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    return final_record


def process_folder(folder: Path, vault: Path = DEFAULT_VAULT, force: bool = False, model_route: str = "auto") -> PaperRecord:
    record = scan_folder(folder, vault=vault)
    status_path = folder / "处理状态.json"
    result_path = folder / "处理结果.md"
    if record.status == "done" and not force:
        if record.note_path and not record.html_path and Path(record.note_path).exists():
            html_dir = vault / "Knowledge/Paper Deep Readings HTML"
            html_dir.mkdir(parents=True, exist_ok=True)
            note_text = clean_model_output(Path(record.note_path).read_text(encoding="utf-8", errors="replace"))
            generated_title = extract_generated_title(note_text, record.title)
            html_path = unique_html_path(html_dir, generated_title, record.doi_safe)
            write_html_output(note_text, html_path, generated_title, vault)
            record.html_path = str(html_path)
            record.message = "已补生成 HTML 文件"
            record.updated_at = now_iso()
            write_json(status_path, asdict(record))
        return record
    if not record.pdf:
        write_json(status_path, {**asdict(record), "status": "needs_retry", "updated_at": now_iso()})
        return record
    if not record.doi:
        write_json(status_path, {**asdict(record), "status": "needs_review", "updated_at": now_iso()})
        return record

    write_json(status_path, {**asdict(record), "status": "processing", "message": "处理中", "updated_at": now_iso()})
    doi_safe = record.doi_safe
    input_dir = vault / "AI Inputs"
    output_dir = vault / "AI Outputs"
    note_dir = vault / "Knowledge/Paper Deep Readings"
    html_dir = vault / "Knowledge/Paper Deep Readings HTML"
    asset_dir = vault / "Assets/Papers" / doi_safe
    for path in [input_dir, output_dir, note_dir, html_dir, asset_dir]:
        path.mkdir(parents=True, exist_ok=True)
    packet_path = input_dir / f"{doi_safe}.paper_text_packet.md"
    log_path = output_dir / f"{doi_safe}.run_log.md"
    try:
        supplements = find_supplements(folder, Path(record.pdf))
        extract_pdf_packet(Path(record.pdf), packet_path, record.title, record.doi, supplements=supplements)
        generated, source_model = generate_deep_reading(vault, packet_path, model_route=model_route)
        generated = clean_model_output(generated)
        images = classify_images(folder)
        copied = copy_assets(images, asset_dir)
        asset_rel = f"Assets/Papers/{doi_safe}"
        note_body, image_review = insert_images(generated, asset_rel, copied)
        note_body = clean_model_output(note_body)
        generated_title = extract_generated_title(note_body, record.title)
        status_data = read_json(status_path)
        existing_note = status_data.get("note_path", "")
        existing_html = status_data.get("html_path", "")
        note_path = unique_note_path(note_dir, generated_title, doi_safe, existing=existing_note)
        html_path = unique_html_path(html_dir, generated_title, doi_safe, existing=existing_html)
        note_text = note_body
        note_path.write_text(note_text, encoding="utf-8")
        write_html_output(note_text, html_path, generated_title, vault)
        final_record = scan_folder(folder, vault=vault)
        final_record.status = "done" if not final_record.missing_figures and not image_review else "needs_review"
        final_record.message = "已生成 Obsidian 笔记" + (
            f"；缺图提示: {', '.join(final_record.missing_figures)}" if final_record.missing_figures else ""
        ) + (
            f"；图片需复核: {'; '.join(image_review)}" if image_review else ""
        )
        final_record.note_path = str(note_path)
        final_record.html_path = str(html_path)
        final_record.supplements = [p.name for p in supplements]
        final_record.image_review = image_review
        final_record.updated_at = now_iso()
        write_json(status_path, asdict(final_record))
        result_path.write_text(
            "\n".join(
                [
                    "# 处理结果",
                    "",
                    f"- 状态：{final_record.status}",
                    f"- DOI：{record.doi}",
                    f"- 标题：{record.title}",
                    f"- 补充材料：{', '.join(final_record.supplements or []) or '无'}",
                    f"- 图片复核：{'; '.join(final_record.image_review or []) or '无'}",
                    f"- Obsidian 笔记：{note_path}",
                    f"- HTML 文件：{html_path}",
                    f"- 文本包：{packet_path}",
                    f"- 图片目录：{asset_dir}",
                    f"- 日志：{log_path}",
                    f"- 更新时间：{final_record.updated_at}",
                ]
            )
            + "\n",
            encoding="utf-8",
        )
        log_path.write_text(
            json.dumps(
                {
                    "record": asdict(final_record),
                    "packet": str(packet_path),
                    "note": str(note_path),
                    "html": str(html_path),
                    "assets": str(asset_dir),
                    "model": source_model,
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        return final_record
    except Exception as exc:
        failed = scan_folder(folder, vault=vault)
        failed.status = "needs_retry"
        error_text = f"{type(exc).__name__}: {exc}" if str(exc) else type(exc).__name__
        failed.message = f"处理失败: {error_text}"
        failed.updated_at = now_iso()
        write_json(status_path, asdict(failed))
        failure_report = "\n".join(
            [
                "# 处理失败",
                "",
                f"- 错误类型：{type(exc).__name__}",
                f"- 错误信息：{str(exc) or '(空错误信息)'}",
                f"- 文本包：{packet_path}",
                f"- 更新时间：{failed.updated_at}",
                "",
                "```text",
                traceback.format_exc().strip(),
                "```",
                "",
            ]
        )
        log_path.write_text(failure_report, encoding="utf-8")
        result_path.write_text(failure_report, encoding="utf-8")
        return failed


def refresh_assets(folder: Path, vault: Path = DEFAULT_VAULT) -> PaperRecord:
    record = scan_folder(folder, vault=vault)
    status_path = folder / "处理状态.json"
    if not record.doi:
        record.status = "needs_review"
        record.message = "无法刷新图片：未识别 DOI"
        write_json(status_path, asdict(record))
        return record
    images = classify_images(folder)
    asset_dir = vault / "Assets/Papers" / record.doi_safe
    copy_assets(images, asset_dir)
    refreshed = scan_folder(folder, vault=vault)
    refreshed.status = "done" if not refreshed.missing_figures and not refreshed.image_review else "needs_review"
    refreshed.message = "图片资产已刷新" + (
        f"；缺图提示: {', '.join(refreshed.missing_figures)}" if refreshed.missing_figures else ""
    )
    refreshed.updated_at = now_iso()
    write_json(status_path, asdict(refreshed))
    return refreshed


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--inbox", type=Path, default=DEFAULT_INBOX)
    parser.add_argument("--vault", type=Path, default=DEFAULT_VAULT)
    parser.add_argument("--scan", action="store_true")
    parser.add_argument("--prepare", type=Path, default=None, help="Prepare packet and auto-extracted PDF figure pages for agent-side generation")
    parser.add_argument("--finalize", type=Path, default=None, help="Finalize one folder using --generated-md from the current agent")
    parser.add_argument("--generated-md", type=Path, default=None, help="Markdown generated by Codex/Claude for --finalize")
    parser.add_argument("--no-auto-images", action="store_true", help="Do not render PDF figure pages during --prepare")
    parser.add_argument("--process", type=Path, default=None, help="Folder name or path to process")
    parser.add_argument("--process-all", action="store_true")
    parser.add_argument("--refresh-assets", type=Path, default=None, help="Refresh copied image assets for one folder")
    parser.add_argument("--force", action="store_true")
    parser.add_argument("--model", choices=["auto", "deepseek", "claude"], default="auto")
    args = parser.parse_args()

    if args.scan or (not args.prepare and not args.finalize and not args.process and not args.process_all and not args.refresh_assets):
        print(json.dumps([asdict(r) for r in scan_inbox(args.inbox, args.vault)], ensure_ascii=False, indent=2))
        return 0
    if args.prepare:
        folder = args.prepare
        if not folder.is_absolute():
            folder = args.inbox / folder
        record = prepare_folder(folder, vault=args.vault, extract_auto_images=not args.no_auto_images)
        print(json.dumps(asdict(record), ensure_ascii=False, indent=2))
        return 0 if record.status in {"prepared", "needs_review", "done"} else 1
    if args.finalize:
        if not args.generated_md:
            parser.error("--finalize requires --generated-md")
        folder = args.finalize
        if not folder.is_absolute():
            folder = args.inbox / folder
        record = finalize_folder(folder, args.generated_md, vault=args.vault)
        print(json.dumps(asdict(record), ensure_ascii=False, indent=2))
        return 0 if record.status in {"done", "needs_review"} else 1
    if args.process:
        folder = args.process
        if not folder.is_absolute():
            folder = args.inbox / folder
        record = process_folder(folder, vault=args.vault, force=args.force, model_route=args.model)
        print(json.dumps(asdict(record), ensure_ascii=False, indent=2))
        return 0 if record.status in {"done", "needs_review"} else 1
    if args.refresh_assets:
        folder = args.refresh_assets
        if not folder.is_absolute():
            folder = args.inbox / folder
        record = refresh_assets(folder, vault=args.vault)
        print(json.dumps(asdict(record), ensure_ascii=False, indent=2))
        return 0 if record.status in {"done", "needs_review"} else 1
    if args.process_all:
        results = []
        for record in scan_inbox(args.inbox, args.vault):
            if args.force or record.status in {"new", "needs_retry"} or (record.status == "done" and record.note_path and not record.html_path):
                results.append(process_folder(Path(record.folder_path), vault=args.vault, force=args.force, model_route=args.model))
            else:
                results.append(record)
        print(json.dumps([asdict(r) for r in results], ensure_ascii=False, indent=2))
        return 0
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
